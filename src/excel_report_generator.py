from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from report_models import MemberAttendance

class ExcelReportGenerator:
    def generate(
            self,
            members: list[MemberAttendance],
            title: str) -> BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Declining Attendance"
        TITLE_ROW = 1
        HEADER_ROW = 3
        TITLE_FONT_SIZE = 16

        headers = [
            "First Name",
            "Last Name",
            "Early Period Attendance",
            "Worship Day Count",
            "Early Period Frequency",
            "Late Period Attendance",
            "Worship Day Count",
            "Late Period Frequency",
            "Frequency Change",
        ]

        title_cell = sheet.cell(row=TITLE_ROW, column=1, value=title)
        sheet.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=len(headers))
        title_cell.alignment = Alignment(horizontal="center")
        title_cell.font = Font(size=TITLE_FONT_SIZE)

        PERCENTAGE_FORMAT = "0%"
        PERCENTAGE_COLUMN_HEADERS = ["Early Period Frequency", "Late Period Frequency", "Frequency Change"]
        percentage_columns = [get_column_letter(headers.index(header) + 1) for header in PERCENTAGE_COLUMN_HEADERS]

        sheet.append([])
        sheet.append(headers)
        sheet.row_dimensions[HEADER_ROW].height = 30
        for column_index, cell in enumerate(sheet[HEADER_ROW], start=1):
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True)
            sheet.column_dimensions[cell.column_letter].width = 16

        for member in members:
            row_index = sheet.max_row + 1
            sheet.append([
                member.first_name,
                member.last_name,
                member.early_period_attendance,
                member.early_period_record_count,
                member.early_period_frequency(),
                member.late_period_attendance,
                member.late_period_record_count,
                member.late_period_frequency(),
                member.frequency_change(),
            ])
            for column_letter in percentage_columns:
                sheet[f"{column_letter}{row_index}"].number_format = PERCENTAGE_FORMAT

        excel_bytes = BytesIO()
        workbook.save(excel_bytes)
        return excel_bytes
